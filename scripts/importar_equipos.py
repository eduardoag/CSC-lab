# ============================================================
# CSC LAB
# Importador de equipos desde Excel
# ============================================================

from pathlib import Path

from openpyxl import load_workbook

from core.equipos import (
    inicializar_equipos,
    obtener_conexion,
)


RUTA_PROYECTO = Path(__file__).resolve().parent.parent

RUTA_EXCEL = (
    RUTA_PROYECTO
    / "data"
    / "importaciones"
    / "grupos_alumnos_csc-lab.xlsx"
)


HOJAS = {
    "grupos_alumnos_3A": "3A",
    "grupos_alumnos_3B": "3B",
    "grupos_alumnos_5A": "5A",
    "grupos_alumnos_5B": "5B",
}


def normalizar_grupo(valor):
    """
    Convierte el número de grupo del Excel a entero.
    Si la celda está vacía, devuelve None.
    """

    if valor is None:
        return None

    if isinstance(valor, str):
        valor = valor.strip()

        if not valor:
            return None

    try:
        return int(valor)

    except (TypeError, ValueError):
        raise ValueError(
            f"Número de grupo inválido: {valor}"
        )


def importar_equipos():

    if not RUTA_EXCEL.exists():
        raise FileNotFoundError(
            f"No se encontró el archivo:\n{RUTA_EXCEL}"
        )

    inicializar_equipos()

    libro = load_workbook(
        RUTA_EXCEL,
        data_only=True,
    )

    equipos_creados = set()
    alumnos_asignados = 0
    alumnos_sin_equipo = 0

    with obtener_conexion() as conexion:

        for nombre_hoja, curso_esperado in HOJAS.items():

            if nombre_hoja not in libro.sheetnames:
                raise ValueError(
                    f"Falta la hoja {nombre_hoja}"
                )

            hoja = libro[nombre_hoja]

            encabezados = [
                celda.value
                for celda in hoja[1]
            ]

            columnas = {
                nombre: indice
                for indice, nombre
                in enumerate(encabezados)
            }

            requeridas = {
                "alumno_id",
                "curso",
                "Grupo",
            }

            faltantes = (
                requeridas - columnas.keys()
            )

            if faltantes:
                raise ValueError(
                    f"{nombre_hoja}: faltan columnas "
                    f"{sorted(faltantes)}"
                )

            print()
            print(
                f"--- {curso_esperado} ---"
            )

            for fila in hoja.iter_rows(
                min_row=2,
                values_only=True,
            ):

                alumno_id = fila[
                    columnas["alumno_id"]
                ]

                curso = fila[
                    columnas["curso"]
                ]

                grupo = normalizar_grupo(
                    fila[columnas["Grupo"]]
                )

                if alumno_id is None:
                    continue

                alumno_id = str(
                    alumno_id
                ).strip()

                curso = str(
                    curso
                ).strip()

                if curso != curso_esperado:
                    raise ValueError(
                        f"{alumno_id}: curso {curso} "
                        f"no coincide con "
                        f"{curso_esperado}"
                    )

                # --------------------------------------------
                # ALUMNO TODAVÍA SIN EQUIPO
                # --------------------------------------------

                if grupo is None:

                    alumnos_sin_equipo += 1

                    print(
                        f"  ⏳ {alumno_id}: "
                        "sin equipo"
                    )

                    continue

                # --------------------------------------------
                # CREAR / ACTUALIZAR EQUIPO
                # --------------------------------------------

                equipo_id = (
                    f"CSC-{curso}-E{grupo:02d}"
                )

                nombre_equipo = (
                    f"Equipo {grupo}"
                )

                conexion.execute(
                    """
                    INSERT INTO equipos (
                        equipo_id,
                        curso,
                        numero_equipo,
                        nombre_equipo,
                        activo
                    )
                    VALUES (?, ?, ?, ?, 1)

                    ON CONFLICT(equipo_id)
                    DO UPDATE SET
                        curso = excluded.curso,
                        numero_equipo =
                            excluded.numero_equipo,
                        nombre_equipo =
                            excluded.nombre_equipo,
                        activo = 1
                    """,
                    (
                        equipo_id,
                        curso,
                        grupo,
                        nombre_equipo,
                    ),
                )

                equipos_creados.add(
                    equipo_id
                )

                # --------------------------------------------
                # CERRAR OTRA MEMBRESÍA ACTIVA
                # --------------------------------------------

                conexion.execute(
                    """
                    UPDATE miembros_equipo
                    SET activo = 0,
                        fecha_hasta =
                            CURRENT_TIMESTAMP
                    WHERE alumno_id = ?
                      AND activo = 1
                      AND equipo_id != ?
                    """,
                    (
                        alumno_id,
                        equipo_id,
                    ),
                )

                # --------------------------------------------
                # COMPROBAR SI YA PERTENECE AL EQUIPO
                # --------------------------------------------

                existente = conexion.execute(
                    """
                    SELECT id
                    FROM miembros_equipo
                    WHERE alumno_id = ?
                      AND equipo_id = ?
                      AND activo = 1
                    """,
                    (
                        alumno_id,
                        equipo_id,
                    ),
                ).fetchone()

                if existente is None:

                    conexion.execute(
                        """
                        INSERT INTO miembros_equipo (
                            equipo_id,
                            alumno_id,
                            activo
                        )
                        VALUES (?, ?, 1)
                        """,
                        (
                            equipo_id,
                            alumno_id,
                        ),
                    )

                alumnos_asignados += 1

        conexion.commit()

    print()
    print("=" * 55)
    print("IMPORTACIÓN FINALIZADA")
    print("=" * 55)

    print(
        f"Equipos detectados: "
        f"{len(equipos_creados)}"
    )

    print(
        f"Alumnos con equipo: "
        f"{alumnos_asignados}"
    )

    print(
        f"Alumnos todavía sin equipo: "
        f"{alumnos_sin_equipo}"
    )


if __name__ == "__main__":
    importar_equipos()
